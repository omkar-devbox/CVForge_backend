"""Unit tests for the new document extraction toolstack:
- PyMuPDF (fitz)
- pdfplumber
- LibreOffice .doc -> .docx
- BeautifulSoup & regex cleaner
- RapidFuzz skill normalization
- dateparser date parsing
- Pydantic schema validation
"""

from pathlib import Path
import tempfile
import docx
import pymupdf as fitz
import pytest

from app.services.file_manifest.extractors import (
    DocExtractor,
    DocxExtractor,
    EntityExtractor,
    PDFExtractor,
)
from app.services.file_manifest.schemas import (
    CandidateLinks,
    EducationItem,
    ExperienceItem,
    ExtractedDocument,
)
from app.services.file_manifest.text_processors import (
    ResumeDateParser,
    SkillNormalizer,
    TextCleaner,
)


class TestTextCleaner:
    def test_clean_html_and_markdown(self):
        dirty = "<p>Experienced <b>Software Engineer</b> with &amp; Python skills.</p>\n<script>alert(1)</script>"
        clean = TextCleaner.clean(dirty)
        assert "<p>" not in clean
        assert "<script>" not in clean
        assert "&amp;" not in clean
        assert "Experienced Software Engineer with & Python skills." in clean

    def test_clean_bullets_and_unicode(self):
        dirty = "• Python\n● FastAPI\n➢ Docker\u200b\nPage 1 of 2"
        clean = TextCleaner.clean(dirty)
        assert "\u200b" not in clean
        assert "Page 1 of 2" not in clean
        assert "- Python" in clean
        assert "- FastAPI" in clean
        assert "- Docker" in clean

    def test_stitch_fragmented_lines_conservative(self):
        # Must NOT stitch title and company
        text = "Python Developer\nABC Technologies"
        assert "Python Developer ABC Technologies" not in TextCleaner.clean(text)

        # Must NOT stitch bullet title and company
        bullet_text = "- Python Developer\nABC Technologies"
        assert "- Python Developer ABC Technologies" not in TextCleaner.clean(bullet_text)

        # Must NOT stitch lines separated by terminal punctuation
        terminal_text = "- Completed project on time.\nStarted new initiative."
        cleaned_terminal = TextCleaner.clean(terminal_text)
        assert "Completed project on time. Started" not in cleaned_terminal

        # Should stitch when next line is a lowercase continuation
        wrapped_bullet = "- Developed scalable services\nfor high throughput data processing."
        cleaned_wrapped = TextCleaner.clean(wrapped_bullet)
        assert "- Developed scalable services for high throughput data processing." in cleaned_wrapped

        # Should stitch when prev line ends with hanging comma or conjunction
        hanging_bullet = "- Designed microservices with Python and\nFastAPI framework."
        cleaned_hanging = TextCleaner.clean(hanging_bullet)
        assert "- Designed microservices with Python and FastAPI framework." in cleaned_hanging

    def test_context_aware_page_removal(self):
        # Standalone page markers should be removed
        standalone = "Experience\nPage 2 of 5\nSenior Engineer\npage 2\n- 2 -"
        cleaned_standalone = TextCleaner.clean(standalone)
        assert "Page 2 of 5" not in cleaned_standalone
        assert "page 2" not in cleaned_standalone
        assert "- 2 -" not in cleaned_standalone
        assert "Experience" in cleaned_standalone
        assert "Senior Engineer" in cleaned_standalone

        # Inline mentions of page in legitimate text must be preserved
        inline_doc = "Published research paper in IEEE, page 2 of volume 4."
        cleaned_inline = TextCleaner.clean(inline_doc)
        assert "page 2 of volume 4" in cleaned_inline


class TestSkillNormalizer:
    def test_canonical_mapping(self):
        assert SkillNormalizer.normalize_single_skill("reactjs") == "React"
        assert SkillNormalizer.normalize_single_skill("React.js") == "React"
        assert SkillNormalizer.normalize_single_skill("postgres") == "PostgreSQL"
        assert SkillNormalizer.normalize_single_skill("k8s") == "Kubernetes"
        assert SkillNormalizer.normalize_single_skill("golang") == "Go"

    def test_deduplication_with_rapidfuzz(self):
        skills = [
            "React",
            "reactjs",
            "React.js",
            "PostgreSQL",
            "postgres",
            "FastAPI",
            "fastapi",
            "Docker",
        ]
        normalized = SkillNormalizer.normalize_and_deduplicate(skills)
        assert "React" in normalized
        # Should not have duplicate react variants
        assert len([s for s in normalized if "React" in s]) == 1
        assert "PostgreSQL" in normalized
        assert len([s for s in normalized if "Postgres" in s or "PostgreSQL" in s]) == 1
        assert "FastAPI" in normalized

    def test_distinct_skills_never_merged(self):
        # Java and JavaScript must NEVER merge
        result_java = SkillNormalizer.normalize_and_deduplicate(["Java", "JavaScript"])
        assert "Java" in result_java
        assert "JavaScript" in result_java

        # C, C++, and C# must NEVER merge, and 1-letter 'C' must not be dropped
        result_c = SkillNormalizer.normalize_and_deduplicate(["C", "C++", "C#", "R"])
        assert "C" in result_c
        assert "C++" in result_c
        assert "C#" in result_c
        assert "R" in result_c

        # React and React Native must NEVER merge
        result_react = SkillNormalizer.normalize_and_deduplicate(["React", "React Native"])
        assert "React" in result_react
        assert "React Native" in result_react

        # SQL and NoSQL must NEVER merge
        result_sql = SkillNormalizer.normalize_and_deduplicate(["SQL", "NoSQL"])
        assert "SQL" in result_sql
        assert "NoSQL" in result_sql


class TestDateParser:
    def test_date_range_parsing(self):
        start, end, is_current, disp = ResumeDateParser.parse_date_range("Jan 2021 - Present")
        assert start == "2021-01"
        assert end is None
        assert is_current is True
        assert disp == "2021-01 - Present"

    def test_past_date_range(self):
        start, end, is_current, disp = ResumeDateParser.parse_date_range("06/2018 to 08/2021")
        assert start == "2018-06"
        assert end == "2021-08"
        assert is_current is False
        assert disp == "2018-06 - 2021-08"

    def test_single_year(self):
        start, end, is_current, disp = ResumeDateParser.parse_date_range("2020")
        assert start is None
        assert end == "2020"
        assert is_current is False
        assert disp == "2020"

    def test_robust_compact_date_formats(self):
        # Two years without spaces
        s, e, c, d = ResumeDateParser.parse_date_range("2020-2022")
        assert s == "2020"
        assert e == "2022"
        assert c is False

        # Month Year without spaces around hyphen
        s, e, c, d = ResumeDateParser.parse_date_range("Jan 2020-Jan 2022")
        assert s == "2020-01"
        assert e == "2022-01"
        assert c is False

        # Glued month and year
        s, e, c, d = ResumeDateParser.parse_date_range("Jan2020-Jan2022")
        assert s == "2020-01"
        assert e == "2022-01"
        assert c is False

    def test_directional_dates(self):
        # 'Since 2021' should be recognized as ongoing with start_date
        s, e, c, d = ResumeDateParser.parse_date_range("Since 2021")
        assert s == "2021"
        assert e is None
        assert c is True

        # 'From Jan 2020'
        s, e, c, d = ResumeDateParser.parse_date_range("From Jan 2020")
        assert s == "2020-01"
        assert e is None
        assert c is True

        # 'Until 2022'
        s, e, c, d = ResumeDateParser.parse_date_range("Until 2022")
        assert s is None
        assert e == "2022"
        assert c is False

    def test_single_date_custom_default(self):
        # Default 'end'
        s, e, c, d = ResumeDateParser.parse_date_range("2022", default_single_as="end")
        assert s is None and e == "2022"

        # Explicit 'start'
        s, e, c, d = ResumeDateParser.parse_date_range("2022", default_single_as="start")
        assert s == "2022" and e is None

        # Explicit 'both'
        s, e, c, d = ResumeDateParser.parse_date_range("2022", default_single_as="both")
        assert s == "2022" and e == "2022"


class TestDocumentExtractors:
    def test_docx_extractor(self, tmp_path):
        docx_file = tmp_path / "Jane_Doe_Resume.docx"
        doc = docx.Document()
        doc.add_heading("Jane Doe", level=0)
        doc.add_paragraph("Email: jane.doe@example.com | Phone: +1 555-123-4567")
        doc.add_paragraph("LinkedIn: https://www.linkedin.com/in/janedoe")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("ReactJS, Node.js, Python, PostgreSQL, AWS")
        doc.add_heading("Experience", level=1)
        doc.add_paragraph("Senior Full Stack Developer")
        doc.add_paragraph("Acme Corp")
        doc.add_paragraph("March 2021 - Present")
        doc.add_paragraph("• Architected cloud-native microservices")
        doc.add_paragraph("• Led a team of 5 engineers")
        doc.add_heading("Education", level=1)
        doc.add_paragraph("B.Tech in Computer Science | Stanford University | 2017 - 2021")
        doc.save(str(docx_file))

        extractor = DocxExtractor(enable_llm=False)
        extracted = extractor.extract(docx_file)

        assert extracted.status == "success"
        assert extracted.name == "Jane Doe"
        assert "jane.doe@example.com" in extracted.email
        assert any("555" in p for p in extracted.contact_no)
        assert extracted.links.linkedin == "https://www.linkedin.com/in/janedoe"
        # Check skill normalization
        assert "React" in extracted.skills
        assert "Node.js" in extracted.skills
        assert "PostgreSQL" in extracted.skills
        # Check parsed dates in experience
        assert len(extracted.experience) >= 1
        exp = extracted.experience[0]
        assert exp.is_current is True
        assert exp.start_date == "2021-03"
        # Check parsed dates in education
        assert len(extracted.education) >= 1
        edu = extracted.education[0]
        assert edu.degree is not None
        assert edu.institution is not None

    def test_pdf_extractor_pymupdf(self, tmp_path):
        pdf_file = tmp_path / "John_Smith_Resume.pdf"
        doc = fitz.open()
        page = doc.new_page()
        text = """John Smith
Email: john.smith@example.com
Phone: +91 9876543210
LinkedIn: https://www.linkedin.com/in/johnsmith

Skills
Python, FastAPI, Docker, Kubernetes

Experience
Backend Engineer
TechCorp India
01/2020 - 12/2022
- Developed high-performance REST APIs

Education
B.E. Mechanical Engineering
Pune University
2015 - 2019
"""
        page.insert_text((50, 72), text, fontsize=11)
        doc.save(str(pdf_file))
        doc.close()

        extractor = PDFExtractor(enable_llm=False)
        extracted = extractor.extract(pdf_file)

        assert extracted.status == "success"
        assert extracted.name == "John Smith"
        assert "john.smith@example.com" in extracted.email
        assert any("9876543210" in p for p in extracted.contact_no)
        assert "Python" in extracted.skills
        assert "FastAPI" in extracted.skills
        assert "Kubernetes" in extracted.skills
        assert len(extracted.experience) >= 1
        assert extracted.experience[0].start_date == "2020-01"
        assert extracted.experience[0].end_date == "2022-12"

    def test_gemma_extractor_offline(self):
        from app.services.file_manifest.ai_models import GemmaExtractor
        gemma = GemmaExtractor()
        assert gemma.is_available() is True
        assert gemma._initialize() is True

    def test_infer_candidate_name_cleaning(self):
        extractor = EntityExtractor(enable_llm=False)
        cases = [
            ("AadeshAbichandani[5y_0m]-WA-npu- call later-26 aug.docx", "Aadesh Abichandani"),
            ("Aakrati Agrawal--npu.docx", "Aakrati Agrawal"),
            ("Abdul Basith P M.DOCX", "Abdul Basith P M"),
            ("Abdul Razzaq.DOCX", "Abdul Razzaq"),
            ("JaneDoe_Resume_2024.docx", "Jane Doe"),
            ("MD AFTAB ALAM--54% in 10th.docx", "Md Aftab Alam"),
            ("Prasad Varne--cant upload-31 mar-8 apr-26 may-9 jun.docx", "Prasad Varne"),
            ("Sourabh_Soni_ask if he has exp in ev powertrain-if yest take for 14410- fwd.docx", "Sourabh Soni"),
            ("Puneeth[5y_0m]-WA-not looking change-25 Aug.pdf", "Puneeth"),
            ("MayureshKhadse[8y_0m]-not suitable JD-26 aug.pdf", "Mayuresh Khadse"),
        ]
        for filename, expected_name in cases:
            inferred = extractor.infer_candidate_name("", file_path=Path(filename))
            assert inferred == expected_name, f"Expected {expected_name}, got {inferred} for {filename}"

    def test_phone_number_extraction_variations(self):
        extractor = EntityExtractor(enable_llm=False)
        text = """
        Candidate Phone Numbers:
        Mobile: +91 98765 43210
        Alt: 98765-43211
        Office: +1 (555) 123-4567
        """
        phones = extractor.extract_phones(text)
        assert len(phones) >= 2
        assert any("98765 43210" in p or "98765" in p for p in phones)

    def test_json_standard_format_no_redundant_entities(self, tmp_path):
        import json
        from app.services.file_manifest.service import FileManifestService
        from app.core.config import settings

        docs_dir = tmp_path / "StandardDocs"
        docs_dir.mkdir()
        docx_file = docs_dir / "StandardCandidate.docx"

        doc = docx.Document()
        doc.add_heading("Standard Candidate", level=0)
        doc.add_paragraph("Email: standard.cand@example.com | Phone: +91 98765 43210")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Python, FastAPI, Docker")
        doc.save(str(docx_file))

        service = FileManifestService(
            config=settings,
            base_path=str(docs_dir),
            extracted_data_folder_name="Extracted data",
            enable_llm=False,
        )

        extracted = service.process_file(docx_file, save_json=True, compute_embedding=False)
        assert extracted.status == "success"
        assert extracted.name == "Standard Candidate"
        assert "standard.cand@example.com" in extracted.email

        # 1. Verify JSON on disk has clean single standard schema
        json_path = Path(extracted.json_output_path)
        assert json_path.exists()
        with open(json_path, "r", encoding="utf-8") as f:
            disk_data = json.load(f)

        # Must have standard top-level fields
        assert "file_name" in disk_data
        assert "status" in disk_data
        assert "name" in disk_data
        assert "contact_no" in disk_data
        assert "email" in disk_data
        assert "skills" in disk_data
        assert "education" in disk_data
        assert "experience" in disk_data
        assert "ats_score" in disk_data

        # Must NOT have redundant internal entities object or heavy vectors
        assert "entities" not in disk_data
        assert "embeddings" not in disk_data
        assert "embedding" not in disk_data

        # 2. Verify loading back from disk populates doc.entities in memory
        loaded_doc = service.get_extracted_document("StandardCandidate")
        assert loaded_doc is not None
        assert "standard.cand@example.com" in loaded_doc.entities.emails
        assert "python" in loaded_doc.entities.skills

    def test_docx_table_project_extraction(self, tmp_path):
        docx_file = tmp_path / "Candidate_With_Tables.docx"
        doc = docx.Document()
        doc.add_heading("Alex Rivera", level=0)
        doc.add_paragraph("Email: alex.r@example.com | Phone: 9876543210")
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Python, FastAPI")
        
        # Add table with projects
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Project"
        table.rows[0].cells[1].text = "Technologies"
        table.rows[1].cells[0].text = "CVForge"
        table.rows[1].cells[1].text = "Python, FastAPI"
        doc.save(str(docx_file))

        extractor = DocxExtractor(enable_llm=False)
        extracted = extractor.extract(docx_file)

        assert extracted.status == "success"
        assert extracted.name == "Alex Rivera"
        # Table headers must not be in skills
        assert "Project" not in extracted.skills
        assert "Technologies" not in extracted.skills
        # Project should be parsed in projects
        assert len(extracted.projects) == 1
        assert extracted.projects[0].name == "CVForge"
        assert "Python" in extracted.projects[0].technologies
        assert "FastAPI" in extracted.projects[0].technologies

    def test_docx_external_relationship_no_crash(self, tmp_path):
        docx_file = tmp_path / "External_Rel_Test.docx"
        doc = docx.Document()
        doc.add_paragraph("Abdul Razzaq")
        doc.add_paragraph("Email: abdul.razzaq@example.com | Phone: 9876543210")
        doc.part.rels.add_relationship(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
            "http://example.com/external_logo.png",
            "rIdExternalLogo",
            is_external=True,
        )
        doc.save(str(docx_file))

        extractor = DocxExtractor(enable_llm=False)
        images_dir = tmp_path / "extracted_images"
        extracted = extractor.extract(docx_file, images_output_dir=images_dir)

        assert extracted.status == "success"
        assert extracted.name == "Abdul Razzaq"

    def test_clear_session_retains_model_weights(self):
        from app.services.file_manifest.ai_models import GemmaExtractor

        gemma = GemmaExtractor()
        if gemma.is_available():
            assert gemma._initialize() is True
            assert gemma._session is not None
            gemma.clear_session()
            # Model weights and session must still remain active in memory
            assert gemma._session is not None
            assert gemma._is_initialized is True



