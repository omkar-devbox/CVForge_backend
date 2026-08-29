"""DOC (legacy Microsoft Word 97-2003 binary) document extractor."""

import logging
from pathlib import Path
import re
import struct
from typing import List, Optional
from app.services.file_manifest.extractors.base import BaseDocumentExtractor
from app.services.file_manifest.extractors.entity_extractor import EntityExtractor
from app.services.file_manifest.schemas import (
    ExtractedDocument,
)

logger = logging.getLogger("cvforge.extractor.doc")


class DocExtractor(BaseDocumentExtractor):
    """Extracts text, metadata, and entities from legacy .doc binary files."""

    def __init__(self):
        self.entity_extractor = EntityExtractor()

    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in {".doc"}

    def extract(
        self, file_path: Path, images_output_dir: Optional[Path] = None
    ) -> ExtractedDocument:
        if not file_path.exists():
            return self.create_failed_document(
                file_path, f"File does not exist: {file_path}"
            )

        try:
            size_bytes, size_kb, mime_type = self.get_file_metadata_stats(file_path)
            sha256 = self.compute_sha256(file_path)
        except Exception as exc:
            return self.create_failed_document(
                file_path, f"Failed to read file stats: {exc}"
            )

        extracted_text: Optional[str] = None
        error_msg: Optional[str] = None

        # Strategy 1: Try reading as DOCX (many modern exports have .doc extension but are docx XML)
        try:
            import docx

            doc = docx.Document(str(file_path))
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if paras:
                extracted_text = "\n\n".join(paras)
        except Exception:
            pass

        # Strategy 2: MS-DOC Piece Table (FIB Clx) extraction via olefile
        if not extracted_text:
            try:
                extracted_text = self._extract_msdoc_piece_table(file_path)
            except Exception as ole_exc:
                logger.debug(
                    f"Piece table extraction failed for {file_path.name}: {ole_exc}"
                )

        # Strategy 3: Binary text stream recovery fallback
        if not extracted_text:
            try:
                with open(file_path, "rb") as f:
                    content = f.read()
                text_chunks = self._extract_text_from_binary_stream(content)
                if text_chunks:
                    extracted_text = "\n\n".join(text_chunks)
            except Exception as bin_exc:
                error_msg = f"Binary fallback failed: {bin_exc}"

        raw_text = extracted_text or ""
        clean_text = self.clean_extracted_text(raw_text)
        paragraphs = self.split_paragraphs(clean_text)
        word_count, char_count = self.compute_word_character_counts(clean_text)

        extracted_data = self.entity_extractor.extract(clean_text, file_path=file_path)

        status = "success"
        if not clean_text.strip():
            status = "failed"
            error_msg = error_msg or "Could not extract readable text from .doc file"
        elif error_msg:
            status = "partial"

        return ExtractedDocument(
            file_name=file_path.name,
            file_stem=file_path.stem,
            status=status,
            name=extracted_data["name"],
            contact_no=extracted_data["contact_no"],
            email=extracted_data["email"],
            links=extracted_data["links"],
            skills=extracted_data["skills"],
            education=extracted_data["education"],
            experience=extracted_data["experience"],
            images=[],
            embedding=None,
            error_message=error_msg,
        )

    def _extract_msdoc_piece_table(self, file_path: Path) -> Optional[str]:
        """Extract text from MS Word binary format using the Clx Piece Table."""
        import olefile

        if not olefile.isOleFile(str(file_path)):
            return None

        with olefile.OleFileIO(str(file_path)) as ole:
            if not ole.exists("WordDocument"):
                return None
            word_data = ole.openstream("WordDocument").read()

            table_name = (
                "1Table"
                if ole.exists("1Table")
                else ("0Table" if ole.exists("0Table") else None)
            )
            if not table_name:
                return None
            table_data = ole.openstream(table_name).read()

            # FIB: fcClx (offset 0x01A2) and lcbClx (offset 0x01A6)
            if len(word_data) < 0x01AA:
                return None

            fcClx, = struct.unpack_from("<I", word_data, 0x01A2)
            lcbClx, = struct.unpack_from("<I", word_data, 0x01A6)

            if fcClx + lcbClx > len(table_data):
                return None

            clx = table_data[fcClx : fcClx + lcbClx]
            pos = 0
            while pos < len(clx):
                clx_type = clx[pos]
                pos += 1
                if clx_type == 1:  # Grpprl
                    if pos + 2 > len(clx):
                        break
                    cb_grpprl, = struct.unpack_from("<H", clx, pos)
                    pos += 2 + cb_grpprl
                elif clx_type == 2:  # Plcfpcd (Piece table)
                    if pos + 4 > len(clx):
                        break
                    lcb, = struct.unpack_from("<I", clx, pos)
                    pos += 4
                    pcd_data = clx[pos : pos + lcb]
                    n = (lcb - 4) // 12
                    if n <= 0 or (4 * (n + 1) + 8 * n) > len(pcd_data):
                        break
                    cps = [
                        struct.unpack_from("<I", pcd_data, i * 4)[0]
                        for i in range(n + 1)
                    ]
                    pcds_offset = 4 * (n + 1)

                    text_pieces = []
                    for i in range(n):
                        cp_len = cps[i + 1] - cps[i]
                        pcd_bytes = pcd_data[
                            pcds_offset + i * 8 : pcds_offset + (i + 1) * 8
                        ]
                        fcValue, = struct.unpack_from("<I", pcd_bytes, 2)
                        fCompressed = (fcValue & 0x40000000) != 0
                        fc = fcValue & (~0x40000000)

                        if fCompressed:
                            actual_offset = fc // 2
                            piece = word_data[
                                actual_offset : actual_offset + cp_len
                            ].decode("latin1", errors="ignore")
                        else:
                            piece = word_data[fc : fc + cp_len * 2].decode(
                                "utf-16le", errors="ignore"
                            )
                        text_pieces.append(piece)
                    return "".join(text_pieces)
        return None

    def _extract_text_from_binary_stream(self, data: bytes) -> List[str]:
        """Extract readable UTF-16LE and ASCII text blocks from binary streams."""
        chunks: List[str] = []

        # 1. UTF-16LE strings (common in MS Word binary)
        try:
            decoded_utf16 = data.decode("utf-16le", errors="ignore")
            clean_utf16 = re.findall(r"[\w\s.,@:;/\-–—\(\)]{15,}", decoded_utf16)
            for c in clean_utf16:
                c_clean = c.strip()
                if c_clean and len(c_clean) >= 20:
                    chunks.append(c_clean)
        except Exception:
            pass

        # 2. ASCII/UTF-8 strings
        if not chunks:
            try:
                decoded_ascii = data.decode("latin1", errors="ignore")
                clean_ascii = re.findall(
                    r"[A-Za-z0-9\s.,@:;/\-–—\(\)]{15,}", decoded_ascii
                )
                for c in clean_ascii:
                    c_clean = c.strip()
                    if c_clean and len(c_clean) >= 20:
                        chunks.append(c_clean)
            except Exception:
                pass

        return chunks
