"""DOCX document extractor using python-docx."""

import logging
from pathlib import Path
from typing import List, Optional
import docx
from app.services.file_manifest.extractors.base import BaseDocumentExtractor
from app.services.file_manifest.extractors.entity_extractor import EntityExtractor
from app.services.file_manifest.schemas import (
    ExtractedDocument,
)

logger = logging.getLogger("cvforge.extractor.docx")


class DocxExtractor(BaseDocumentExtractor):
    """Extracts text, paragraphs, tables, metadata, and entities from DOCX files."""

    def __init__(self):
        self.entity_extractor = EntityExtractor()

    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in {".docx"}

    def extract(
        self, file_path: Path, images_output_dir: Optional[Path] = None
    ) -> ExtractedDocument:
        if not file_path.exists():
            return self.create_failed_document(
                file_path, f"File does not exist: {file_path}"
            )

        try:
            size_bytes, size_kb, mime_type = self.get_file_metadata_stats(file_path)
            sha256 = self.compute_sha256(file_path)
        except Exception as exc:
            return self.create_failed_document(
                file_path, f"Failed to read file stats: {exc}"
            )

        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            logger.error(f"Failed to open docx file {file_path.name}: {exc}")
            return self.create_failed_document(
                file_path, f"Failed to parse DOCX file: {exc}"
            )

        # 1. Paragraphs & Raw Text
        paragraph_texts: List[str] = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraph_texts.append(text)

        # 2. Tables text extraction
        table_text_lines: List[str] = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                cleaned_cells = []
                last_cell = None
                for c in row_cells:
                    if c != last_cell:
                        cleaned_cells.append(c)
                        last_cell = c
                if any(cleaned_cells):
                    table_text_lines.append(" | ".join(filter(None, cleaned_cells)))

        # Combined text
        all_text_blocks = paragraph_texts + table_text_lines
        raw_text = "\n\n".join(all_text_blocks)
        clean_text = self.clean_extracted_text(raw_text)

        # 4. Extract Images from DOCX
        extracted_images: List[str] = []
        if images_output_dir:
            try:
                images_output_dir.mkdir(parents=True, exist_ok=True)
                img_idx = 1
                for rel in doc.part.rels.values():
                    if "image" in rel.target_ref:
                        img_part = rel.target_part
                        img_blob = img_part.blob
                        # Filter out tiny decorator icons (keep images > 2KB)
                        if len(img_blob) > 2048:
                            ext = Path(rel.target_ref).suffix or ".png"
                            img_filename = "profile_pic" + ext if img_idx == 1 else f"image_{img_idx}{ext}"
                            saved_img_path = images_output_dir / img_filename
                            with open(saved_img_path, "wb") as img_file:
                                img_file.write(img_blob)
                            extracted_images.append(str(saved_img_path.resolve()))
                            img_idx += 1
            except Exception as img_exc:
                logger.warning(f"Image extraction warning for {file_path.name}: {img_exc}")

        # 5. Extract Structured Fields
        extracted_data = self.entity_extractor.extract(clean_text, file_path=file_path)

        return ExtractedDocument(
            file_name=file_path.name,
            file_stem=file_path.stem,
            status="success",
            name=extracted_data["name"],
            contact_no=extracted_data["contact_no"],
            email=extracted_data["email"],
            links=extracted_data["links"],
            skills=extracted_data["skills"],
            education=extracted_data["education"],
            experience=extracted_data["experience"],
            images=extracted_images,
            embedding=None,
            error_message=None,
        )
