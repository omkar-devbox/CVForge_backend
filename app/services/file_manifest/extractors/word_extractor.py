# //------------------------------------------------------------
# // Imports & Dependencies (All Top Side)
# //------------------------------------------------------------
import logging
from pathlib import Path
import re
import shutil
import struct
import subprocess
import tempfile
from typing import Any, List, Optional, Union

try:
    import olefile
except ImportError:
    olefile = None

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.file_manifest.extractors.base import BaseDocumentExtractor
from app.services.file_manifest.extractors.entity_extractor import EntityExtractor
from app.services.file_manifest.schemas import ExtractedDocument, ExtractedTable
from app.services.file_manifest.text_processors import TextCleaner

logger = logging.getLogger("cvforge.extractor.word")


# //------------------------------------------------------------
# // 1. DocxExtractor Implementation (Modern Word OpenXML)
# //------------------------------------------------------------
class DocxExtractor(BaseDocumentExtractor):
    """Extracts text, paragraphs, tables, metadata, and entities from DOCX files."""

    def __init__(
        self,
        enable_gemma: bool = True,
        enable_llm: Optional[bool] = None,
        gemma_model_dir: Optional[Union[str, Path]] = None,
        **kwargs: Any,
    ):
        use_llm = enable_gemma if enable_llm is None else enable_llm
        self.entity_extractor = EntityExtractor(
            enable_llm=use_llm,
            gemma_model_dir=gemma_model_dir,
        )

    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in {".docx"}

    def extract(
        self, file_path: Path, images_output_dir: Optional[Path] = None
    ) -> ExtractedDocument:
        if not file_path.exists():
            return self.create_failed_document(
                file_path, f"File does not exist: {file_path}"
            )

        try:
            size_bytes, size_kb, mime_type = self.get_file_metadata_stats(file_path)
            sha256 = self.compute_sha256(file_path)
        except Exception as exc:
            return self.create_failed_document(
                file_path, f"Failed to read file stats: {exc}"
            )

        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            logger.error(f"Failed to open docx file {file_path.name}: {exc}")
            return self.create_failed_document(
                file_path, f"Failed to parse DOCX file: {exc}"
            )

        # 1. Interleaved Body Extraction (Paragraphs & Tables in true document layout order)
        ordered_text_blocks: List[str] = []
        extracted_tables: List[ExtractedTable] = []

        for child in doc.element.body:
            if child.tag.endswith("p"):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if text:
                    ordered_text_blocks.append(text)
            elif child.tag.endswith("tbl"):
                table = Table(child, doc)
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    cleaned_cells = []
                    last_cell = None
                    for c in row_cells:
                        if c != last_cell:
                            cleaned_cells.append(c)
                            last_cell = c
                    if any(cleaned_cells):
                        row_line = " | ".join(filter(None, cleaned_cells))
                        ordered_text_blocks.append(row_line)
                        table_rows.append(cleaned_cells)
                if table_rows:
                    extracted_tables.append(ExtractedTable(rows=table_rows))

        # Fallback if body iteration produced nothing (e.g. non-standard XML)
        if not ordered_text_blocks:
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    ordered_text_blocks.append(text)

        raw_text = "\n\n".join(ordered_text_blocks)
        clean_text = TextCleaner.clean(raw_text)

        # 2. Extract Images from DOCX (safely skipping external relationships)
        extracted_images: List[str] = []
        if images_output_dir:
            try:
                images_output_dir.mkdir(parents=True, exist_ok=True)
                img_idx = 1
                for rel in doc.part.rels.values():
                    if getattr(rel, "is_external", False):
                        continue
                    try:
                        target_ref = str(getattr(rel, "target_ref", ""))
                        if "image" in target_ref.lower():
                            img_part = getattr(rel, "target_part", None)
                            if not img_part:
                                continue
                            img_blob = getattr(img_part, "blob", None)
                            if img_blob and len(img_blob) > 2048:
                                ext = Path(target_ref).suffix or ".png"
                                if not ext.startswith("."):
                                    ext = f".{ext}"
                                img_filename = "profile_pic" + ext if img_idx == 1 else f"image_{img_idx}{ext}"
                                saved_img_path = images_output_dir / img_filename
                                with open(saved_img_path, "wb") as img_file:
                                    img_file.write(img_blob)
                                extracted_images.append(str(saved_img_path.resolve()))
                                img_idx += 1
                    except Exception as inner_img_exc:
                        logger.debug(f"Skipping individual image in {file_path.name}: {inner_img_exc}")
            except Exception as img_exc:
                logger.warning(f"Image extraction warning for {file_path.name}: {img_exc}")

        # 3. Extract Structured Fields
        extracted_data = self.entity_extractor.extract(
            clean_text, file_path=file_path, tables=extracted_tables
        )

        return ExtractedDocument(
            file_name=file_path.name,
            file_stem=file_path.stem,
            file_size_bytes=size_bytes,
            file_size_kb=size_kb,
            status="success" if clean_text.strip() else "failed",
            name=extracted_data["name"],
            role=extracted_data.get("role"),
            contact_no=extracted_data["contact_no"],
            email=extracted_data["email"],
            links=extracted_data["links"],
            summary=extracted_data.get("summary"),
            skills=extracted_data["skills"],
            education=extracted_data["education"],
            experience=extracted_data["experience"],
            projects=extracted_data.get("projects", []),
            certifications=extracted_data.get("certifications", []),
            languages=extracted_data.get("languages", []),
            tools=extracted_data.get("tools", []),
            sections=extracted_data.get("_sections", {}),
            tables=extracted_tables,
            images=extracted_images,
            embedding=None,
            error_message=None if clean_text.strip() else "No text extracted from DOCX",
        )


# //------------------------------------------------------------
# // 2. DocExtractor Implementation (Legacy Word 97-2003 Binary)
# //------------------------------------------------------------
class DocExtractor(BaseDocumentExtractor):
    """Extracts text, tables, and entities from legacy .doc files using LibreOffice conversion."""

    def __init__(
        self,
        enable_gemma: bool = True,
        enable_llm: Optional[bool] = None,
        gemma_model_dir: Optional[Union[str, Path]] = None,
        **kwargs: Any,
    ):
        use_llm = enable_gemma if enable_llm is None else enable_llm
        self.entity_extractor = EntityExtractor(
            enable_llm=use_llm,
            gemma_model_dir=gemma_model_dir,
        )
        self.docx_extractor = DocxExtractor(
            enable_llm=use_llm,
            gemma_model_dir=gemma_model_dir,
        )

    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in {".doc"}

    def extract(
        self, file_path: Path, images_output_dir: Optional[Path] = None
    ) -> ExtractedDocument:
        if not file_path.exists():
            return self.create_failed_document(
                file_path, f"File does not exist: {file_path}"
            )

        try:
            size_bytes, size_kb, mime_type = self.get_file_metadata_stats(file_path)
            sha256 = self.compute_sha256(file_path)
        except Exception as exc:
            return self.create_failed_document(
                file_path, f"Failed to read file stats: {exc}"
            )

        # -------------------------------------------------------------
        # Strategy 1: Direct DOCX check (many .doc files are modern docx renamed)
        # -------------------------------------------------------------
        try:
            _ = docx.Document(str(file_path))
            logger.info(f"File {file_path.name} is actually DOCX format. Delegating to DocxExtractor.")
            return self.docx_extractor.extract(file_path, images_output_dir=images_output_dir)
        except Exception:
            pass

        # -------------------------------------------------------------
        # Strategy 2: Convert legacy .doc to .docx using headless LibreOffice
        # -------------------------------------------------------------
        soffice_path = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice_path:
            with tempfile.TemporaryDirectory() as temp_dir:
                temp_path = Path(temp_dir)
                try:
                    cmd = [
                        soffice_path,
                        "--headless",
                        "--convert-to",
                        "docx",
                        "--outdir",
                        str(temp_path),
                        str(file_path.resolve()),
                    ]
                    result = subprocess.run(
                        cmd,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.PIPE,
                        timeout=30,
                    )
                    converted_docx = temp_path / f"{file_path.stem}.docx"
                    if converted_docx.exists() and converted_docx.stat().st_size > 0:
                        logger.info(
                            f"Successfully converted {file_path.name} to DOCX using LibreOffice."
                        )
                        doc = self.docx_extractor.extract(
                            converted_docx, images_output_dir=images_output_dir
                        )
                        doc.file_name = file_path.name
                        doc.file_stem = file_path.stem
                        return doc
                    else:
                        logger.warning(
                            f"LibreOffice conversion produced no output for {file_path.name}: {result.stderr.decode('utf-8', errors='ignore')}"
                        )
                except Exception as lo_exc:
                    logger.warning(f"LibreOffice conversion failed for {file_path.name}: {lo_exc}")

        # -------------------------------------------------------------
        # Strategy 3: Fallback Piece Table / Binary Extraction
        # -------------------------------------------------------------
        logger.info(f"Falling back to piece table / binary stream recovery for {file_path.name}")
        extracted_text = None
        error_msg = None

        try:
            extracted_text = self._extract_msdoc_piece_table(file_path)
        except Exception as piece_exc:
            logger.debug(f"Piece table extraction failed for {file_path.name}: {piece_exc}")

        if not extracted_text:
            try:
                with open(file_path, "rb") as f:
                    content = f.read()
                chunks = self._extract_text_from_binary_stream(content)
                if chunks:
                    extracted_text = "\n\n".join(chunks)
            except Exception as bin_exc:
                error_msg = f"Binary fallback failed: {bin_exc}"

        if not extracted_text:
            return self.create_failed_document(
                file_path, error_msg or "Failed to extract text from legacy .doc file"
            )

        clean_text = TextCleaner.clean(extracted_text)
        extracted_data = self.entity_extractor.extract(clean_text, file_path=file_path)

        return ExtractedDocument(
            file_name=file_path.name,
            file_stem=file_path.stem,
            file_size_bytes=size_bytes,
            file_size_kb=size_kb,
            status="success" if clean_text.strip() else "failed",
            name=extracted_data["name"],
            role=extracted_data.get("role"),
            contact_no=extracted_data["contact_no"],
            email=extracted_data["email"],
            links=extracted_data["links"],
            summary=extracted_data.get("summary"),
            skills=extracted_data["skills"],
            education=extracted_data["education"],
            experience=extracted_data["experience"],
            projects=extracted_data.get("projects", []),
            certifications=extracted_data.get("certifications", []),
            languages=extracted_data.get("languages", []),
            tools=extracted_data.get("tools", []),
            images=[],
            embedding=None,
            error_message=error_msg,
        )

    def _extract_msdoc_piece_table(self, file_path: Path) -> Optional[str]:
        """Extracts text from WordDocument stream via OLE Compound File piece table."""
        try:
            if olefile is None or not olefile.isOleFile(str(file_path)):
                return None

            with olefile.OleFileIO(str(file_path)) as ole:
                if not ole.exists("WordDocument"):
                    return None

                word_doc = ole.openstream("WordDocument").read()
                if len(word_doc) < 512:
                    return None

                fib_flags = struct.unpack_from("<H", word_doc, 0x000A)[0]
                table_stream_name = "1Table" if (fib_flags & 0x0200) else "0Table"

                if not ole.exists(table_stream_name):
                    return None

                table_stream = ole.openstream(table_stream_name).read()
                fc_clx = struct.unpack_from("<I", word_doc, 0x01A2)[0]
                lcb_clx = struct.unpack_from("<I", word_doc, 0x01A6)[0]

                if fc_clx + lcb_clx > len(table_stream) or lcb_clx == 0:
                    return None

                clx = table_stream[fc_clx : fc_clx + lcb_clx]
                offset = 0
                while offset < len(clx) and clx[offset] == 1:
                    offset += 1 + 4

                if offset >= len(clx) or clx[offset] != 2:
                    return None

                offset += 1
                lcb_pcdt = struct.unpack_from("<I", clx, offset)[0]
                offset += 4

                pcd_data = clx[offset : offset + lcb_pcdt]
                num_pieces = (len(pcd_data) - 4) // 12
                if num_pieces <= 0:
                    return None

                cp_array = struct.unpack_from(f"<{num_pieces + 1}I", pcd_data, 0)
                pcd_offset = (num_pieces + 1) * 4

                text_parts = []
                for i in range(num_pieces):
                    cp_start = cp_array[i]
                    cp_end = cp_array[i + 1]
                    cp_len = cp_end - cp_start

                    if pcd_offset + 8 > len(pcd_data):
                        break

                    fc_val = struct.unpack_from("<I", pcd_data, pcd_offset + 2)[0]
                    pcd_offset += 8

                    is_unicode = not bool(fc_val & 0x40000000)
                    actual_fc = (fc_val & 0x3FFFFFFF) >> 1 if not is_unicode else (fc_val & 0x3FFFFFFF)

                    if actual_fc >= len(word_doc):
                        continue

                    if is_unicode:
                        byte_len = cp_len * 2
                        chunk = word_doc[actual_fc : actual_fc + byte_len]
                        try:
                            text_parts.append(chunk.decode("utf-16-le", errors="ignore"))
                        except Exception:
                            pass
                    else:
                        chunk = word_doc[actual_fc : actual_fc + cp_len]
                        try:
                            text_parts.append(chunk.decode("cp1252", errors="ignore"))
                        except Exception:
                            pass

                full_text = "".join(text_parts)
                full_text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", full_text)
                return full_text if len(full_text.strip()) > 30 else None

        except Exception as exc:
            logger.debug(f"Piece table extraction failed: {exc}")
            return None

    def _extract_text_from_binary_stream(self, content: bytes) -> List[str]:
        """Recovers clean text strings from binary stream."""
        min_len = 4
        ascii_strings = re.findall(rb"[\x20-\x7E]{" + str(min_len).encode() + rb",}", content)
        chunks = []
        for raw in ascii_strings:
            try:
                decoded = raw.decode("latin-1", errors="ignore").strip()
                if len(decoded) > 10 and any(c.isalpha() for c in decoded):
                    chunks.append(decoded)
            except Exception:
                continue
        return chunks


__all__ = ["DocxExtractor", "DocExtractor"]
